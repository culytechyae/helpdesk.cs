"""Bulk import of users and historical tickets from CSV / XLSX files."""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import from_excel

from extensions import db
from models import (
    User, Ticket, Comment, UserActivityLog, School, ROLE_SUPER_ADMIN, ROLE_ADMIN, ROLE_USER,
    AGENT_ROLES, CATEGORY_AGENT_ROLE, PRIORITIES, STATUSES, utc_now,
)
from user_service import (
    EMAIL_RE, USERNAME_RE, MIN_PASSWORD_LENGTH, assignable_roles, build_user, generate_password,
    make_unique_username, normalize_email,
)

MAX_ROWS = 5000
MAX_FILE_BYTES = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = ('.csv', '.xlsx')


class ImportFileError(Exception):
    """The upload itself is unusable (wrong type, missing columns, too large...)."""


@dataclass
class ImportResult:
    kind: str
    dry_run: bool = False
    total_rows: int = 0
    created: list = field(default_factory=list)     # dicts describing created rows
    skipped: list = field(default_factory=list)     # (row, reason) - e.g. duplicates
    errors: list = field(default_factory=list)      # (row, reason)
    new_users: list = field(default_factory=list)   # users auto-created by a ticket import

    @property
    def ok_count(self):
        return len(self.created)


# --------------------------------------------------------------------------- parsing
def _norm_header(h):
    return re.sub(r'[^a-z0-9]+', '_', str(h or '').strip().lower()).strip('_')


USER_ALIASES = {
    'username': 'username', 'user_name': 'username', 'login': 'username', 'user': 'username',
    'email': 'email', 'e_mail': 'email', 'email_address': 'email', 'mail': 'email',
    'password': 'password', 'temp_password': 'password',
    'role': 'role', 'user_type': 'role', 'type': 'role',
    'full_name': 'full_name', 'name': 'full_name', 'fullname': 'full_name', 'display_name': 'full_name',
    'department': 'department', 'dept': 'department',
    'school': 'school', 'school_slug': 'school', 'school_name': 'school', 'school_code': 'school',
}

TICKET_ALIASES = {
    'title': 'title', 'subject': 'title', 'summary': 'title',
    'description': 'description', 'details': 'description', 'body': 'description',
    'category': 'category', 'department': 'category', 'type': 'category',
    'priority': 'priority', 'severity': 'priority',
    'status': 'status', 'state': 'status',
    'requester': 'requester', 'created_by': 'requester', 'creator': 'requester', 'reporter': 'requester',
    'requester_email': 'requester', 'requested_by': 'requester', 'user': 'requester',
    'assigned_to': 'assignee', 'assignee': 'assignee', 'agent': 'assignee', 'assigned_agent': 'assignee',
    'created_at': 'created_at', 'created': 'created_at', 'created_date': 'created_at', 'date': 'created_at',
    'opened_at': 'created_at',
    'updated_at': 'updated_at', 'updated': 'updated_at', 'closed_at': 'updated_at',
    'resolved_at': 'updated_at', 'closed_date': 'updated_at', 'updated_date': 'updated_at',
    'external_ref': 'external_ref', 'ticket_id': 'external_ref', 'id': 'external_ref',
    'legacy_id': 'external_ref', 'reference': 'external_ref', 'ref': 'external_ref',
    'comments': 'comments', 'comment': 'comments', 'resolution': 'comments', 'notes': 'comments',
    'school': 'school', 'school_slug': 'school', 'school_name': 'school', 'school_code': 'school',
}


def _decode(raw):
    for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportFileError('Could not read the file encoding. Save it as UTF-8 CSV or use .xlsx.')


def read_table(file_storage, aliases, required):
    """Return a list of (row_number, {canonical_field: value}). Row numbers match the sheet."""
    if not file_storage or not file_storage.filename:
        raise ImportFileError('Please choose a file to upload.')
    name = file_storage.filename.lower()
    if not name.endswith(ALLOWED_EXTENSIONS):
        raise ImportFileError('Unsupported file type. Upload a .csv or .xlsx file.')
    raw = file_storage.read(MAX_FILE_BYTES + 1)
    if len(raw) > MAX_FILE_BYTES:
        raise ImportFileError(f'File is too large (limit {MAX_FILE_BYTES // (1024 * 1024)} MB).')
    if not raw.strip():
        raise ImportFileError('The file is empty.')

    if name.endswith('.xlsx'):
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet_rows = list(wb.worksheets[0].iter_rows(values_only=True))
        except Exception:
            raise ImportFileError('Could not read the Excel file. Make sure it is a valid .xlsx workbook.')
    else:
        text = _decode(raw)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        sheet_rows = list(csv.reader(io.StringIO(text), dialect))

    # first non-empty row is the header
    header_idx = next((i for i, r in enumerate(sheet_rows) if any(str(c or '').strip() for c in r)), None)
    if header_idx is None:
        raise ImportFileError('The file has no header row.')
    headers = [aliases.get(_norm_header(h)) for h in sheet_rows[header_idx]]
    missing = [r for r in required if r not in headers]
    if missing:
        raise ImportFileError('Missing required column(s): ' + ', '.join(missing))

    data = []
    for offset, row in enumerate(sheet_rows[header_idx + 1:], start=header_idx + 2):
        if not row or not any(str(c).strip() for c in row if c is not None):
            continue
        record = {}
        for h, cell in zip(headers, row):
            if h and h not in record:
                record[h] = cell.strip() if isinstance(cell, str) else cell
        data.append((offset, record))
    if not data:
        raise ImportFileError('The file contains a header but no data rows.')
    if len(data) > MAX_ROWS:
        raise ImportFileError(f'Too many rows ({len(data)}). Maximum per import is {MAX_ROWS}; split the file.')
    return data


def _s(value):
    """Cell -> stripped string ('' for None)."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


_DATE_FORMATS = (
    '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%dT%H:%M:%S', '%Y-%m-%dT%H:%M', '%Y-%m-%d',
    '%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y', '%d-%m-%Y %H:%M:%S', '%d-%m-%Y %H:%M',
    '%d-%m-%Y', '%d.%m.%Y %H:%M', '%d.%m.%Y',
)


def parse_datetime(value):
    """Parse a spreadsheet date. Text dates are day-first (31/12/2025). Returns naive UTC or None."""
    if value is None or _s(value) == '':
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if hasattr(value, 'year') and hasattr(value, 'month') and not isinstance(value, str):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)):
        try:
            return from_excel(value)
        except Exception:
            raise ValueError(f'Invalid date "{value}"')
    text = str(value).strip()
    text = re.sub(r'\.\d+$', '', text.replace('Z', ''))
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f'Unrecognised date "{value}" (use YYYY-MM-DD or DD/MM/YYYY)')


# --------------------------------------------------------------------------- shared lookups
def _school_lookup():
    schools = School.query.all()
    by_key = {}
    for s in schools:
        by_key[str(s.id)] = s
        by_key[s.slug.lower()] = s
        by_key[s.name.lower()] = s
        if s.display_name:
            by_key[s.display_name.lower()] = s
    return by_key


def _resolve_row_school(actor, row, default_school, schools):
    """School for this row (None, error). School admins are always pinned to their own school."""
    cell = _s(row.get('school'))
    if actor.role != ROLE_SUPER_ADMIN:
        if cell and cell.lower() not in (default_school.slug.lower(), default_school.name.lower(), str(default_school.id)):
            return None, f'You can only import into your own school ("{default_school.name}").'
        return default_school, None
    if cell:
        school = schools.get(cell.lower())
        if not school:
            return None, f'Unknown school "{cell}".'
        return school, None
    if default_school:
        return default_school, None
    return None, 'No school specified (add a "school" column or choose a school).'


ROLE_ALIASES = {
    '': ROLE_USER, 'user': ROLE_USER, 'requester': ROLE_USER, 'staff': ROLE_USER, 'teacher': ROLE_USER,
    'student': ROLE_USER, 'regular user': ROLE_USER, 'regular_user': ROLE_USER,
    'it_agent': 'it_agent', 'it': 'it_agent', 'it agent': 'it_agent',
    'fm_agent': 'fm_agent', 'fm': 'fm_agent', 'fm agent': 'fm_agent', 'facilities': 'fm_agent',
    'hse_agent': 'hse_agent', 'hse': 'hse_agent', 'hse agent': 'hse_agent',
    'admin': ROLE_ADMIN, 'school_admin': ROLE_ADMIN, 'school admin': ROLE_ADMIN,
    'super_admin': ROLE_SUPER_ADMIN, 'super admin': ROLE_SUPER_ADMIN,
}


# --------------------------------------------------------------------------- user import
def import_users(actor, file_storage, default_school, *, dry_run=False, require_password_change=True):
    rows = read_table(file_storage, USER_ALIASES, required=['email'])
    result = ImportResult('users', dry_run=dry_run, total_rows=len(rows))
    schools = _school_lookup() if actor.role == ROLE_SUPER_ADMIN else {}

    taken_users, taken_emails = set(), set()
    for uname, email in db.session.query(User.username, User.email).execution_options(tenant_bypass=True):
        taken_users.add(uname.lower())
        taken_emails.add((email or '').lower())

    for row_no, row in rows:
        try:
            school, err = _resolve_row_school(actor, row, default_school, schools)
            if err:
                result.errors.append((row_no, err))
                continue
            email = normalize_email(_s(row.get('email')))
            if not email or not EMAIL_RE.match(email) or len(email) > 120:
                result.errors.append((row_no, f'Invalid email "{_s(row.get("email"))}".'))
                continue
            role_raw = _s(row.get('role')).lower()
            role = ROLE_ALIASES.get(role_raw)
            if role is None or role not in assignable_roles(actor):
                result.errors.append((row_no, f'Invalid role "{_s(row.get("role"))}". Use: user, it_agent, fm_agent, hse_agent, admin.'))
                continue
            if role == ROLE_SUPER_ADMIN:
                result.errors.append((row_no, 'Super admins cannot be bulk imported.'))
                continue
            if school is None or not school.is_active:
                result.errors.append((row_no, 'School is missing or suspended.'))
                continue

            username = _s(row.get('username'))
            if username:
                if not USERNAME_RE.match(username):
                    result.errors.append((row_no, f'Invalid username "{username}" (3-80 chars; letters, digits . _ @ + -).'))
                    continue
            else:
                username = make_unique_username(email.split('@')[0], school, reserved=taken_users)

            if email in taken_emails:
                result.skipped.append((row_no, f'Email "{email}" already exists.'))
                continue
            if username.lower() in taken_users:
                result.errors.append((row_no, f'Username "{username}" is already taken.'))
                continue

            password = _s(row.get('password'))
            generated = False
            if password:
                if len(password) < MIN_PASSWORD_LENGTH:
                    result.errors.append((row_no, f'Password must be at least {MIN_PASSWORD_LENGTH} characters.'))
                    continue
            else:
                password, generated = generate_password(), True

            user = build_user(
                username=username, email=email, password=password, role=role, school=school,
                full_name=_s(row.get('full_name'))[:120], department=_s(row.get('department'))[:200],
                must_change_password=require_password_change or generated)
            db.session.add(user)
            taken_users.add(username.lower())
            taken_emails.add(email)
            result.created.append({
                'row': row_no, 'username': username, 'email': email, 'role': role,
                'school': school.name, 'password': password if generated else '',
            })
        except Exception as exc:  # one bad row must not abort the file
            result.errors.append((row_no, f'Unexpected error: {exc}'))

    _finish(result, actor, default_school, 'bulk_import_users',
            f'Imported {len(result.created)} users ({len(result.errors)} errors, {len(result.skipped)} skipped)')
    return result


# --------------------------------------------------------------------------- ticket import
CATEGORY_ALIASES = {
    'it': 'IT', 'information technology': 'IT', 'technology': 'IT', 'it support': 'IT',
    'fm': 'FM', 'facilities': 'FM', 'facility': 'FM', 'facilities management': 'FM',
    'maintenance': 'FM', 'handyman': 'FM',
    'hse': 'HSE', 'health and safety': 'HSE', 'health & safety': 'HSE', 'safety': 'HSE',
}
PRIORITY_ALIASES = {
    'low': 'Low', 'minor': 'Low', 'medium': 'Medium', 'normal': 'Medium', 'moderate': 'Medium', '': 'Medium',
    'high': 'High', 'major': 'High', 'critical': 'Critical', 'urgent': 'Critical', 'blocker': 'Critical',
}
STATUS_ALIASES = {
    'open': 'Open', 'new': 'Open', 'pending': 'Open', '': 'Open',
    'assigned': 'Assigned',
    'in progress': 'In Progress', 'in_progress': 'In Progress', 'inprogress': 'In Progress',
    'in-progress': 'In Progress', 'working': 'In Progress', 'on hold': 'In Progress',
    'resolved': 'Resolved', 'fixed': 'Resolved',
    'closed': 'Closed', 'done': 'Closed', 'complete': 'Closed', 'completed': 'Closed', 'cancelled': 'Closed',
}


def import_tickets(actor, file_storage, default_school, *, dry_run=False,
                   create_missing_requesters=False, fallback_to_importer=False):
    rows = read_table(file_storage, TICKET_ALIASES, required=['title'])
    result = ImportResult('tickets', dry_run=dry_run, total_rows=len(rows))
    schools = _school_lookup() if actor.role == ROLE_SUPER_ADMIN else {}
    caches = {}   # school_id -> per-school lookup tables
    now = utc_now().replace(tzinfo=None)
    taken_users, taken_emails = set(), set()
    if create_missing_requesters:
        for uname, email in db.session.query(User.username, User.email).execution_options(tenant_bypass=True):
            taken_users.add(uname.lower())
            taken_emails.add((email or '').lower())

    def cache_for(school):
        if school.id in caches:
            return caches[school.id]
        users = User.query.execution_options(tenant_bypass=True).filter(User.school_id == school.id).all()
        by_ident = {}
        for u in users:
            by_ident[u.username.lower()] = u
            by_ident[(u.email or '').lower()] = u
        refs, keys = set(), set()
        for ref, title, creator, created in db.session.query(
                Ticket.external_ref, Ticket.title, Ticket.created_by, Ticket.created_at
        ).execution_options(tenant_bypass=True).filter(Ticket.school_id == school.id):
            if ref:
                refs.add(ref.lower())
            keys.add((title.lower(), creator, created.replace(microsecond=0) if created else None))
        caches[school.id] = {'users': by_ident, 'refs': refs, 'keys': keys}
        return caches[school.id]

    for row_no, row in rows:
        try:
            school, err = _resolve_row_school(actor, row, default_school, schools)
            if err:
                result.errors.append((row_no, err)); continue
            if school is None or not school.is_active:
                result.errors.append((row_no, 'School is missing or suspended.')); continue
            cache = cache_for(school)

            title = _s(row.get('title'))
            if not title:
                result.errors.append((row_no, 'Title is required.')); continue
            if len(title) > 200:
                result.errors.append((row_no, 'Title is longer than 200 characters.')); continue
            description = _s(row.get('description')) or title

            category = CATEGORY_ALIASES.get(_s(row.get('category')).lower())
            if not category:
                result.errors.append((row_no, f'Invalid category "{_s(row.get("category"))}". Use IT, FM or HSE.')); continue
            priority = PRIORITY_ALIASES.get(_s(row.get('priority')).lower())
            if not priority:
                result.errors.append((row_no, f'Invalid priority "{_s(row.get("priority"))}".')); continue
            status = STATUS_ALIASES.get(_s(row.get('status')).lower())
            if not status:
                result.errors.append((row_no, f'Invalid status "{_s(row.get("status"))}".')); continue

            try:
                created_at = parse_datetime(row.get('created_at')) or now
                updated_at = parse_datetime(row.get('updated_at')) or created_at
            except ValueError as exc:
                result.errors.append((row_no, str(exc))); continue
            if created_at > now + timedelta(days=1):
                result.errors.append((row_no, 'created_at is in the future.')); continue
            if updated_at < created_at:
                updated_at = created_at

            # requester
            requester_ident = _s(row.get('requester')).lower()
            creator = cache['users'].get(requester_ident) if requester_ident else None
            if creator is None:
                if requester_ident and create_missing_requesters and EMAIL_RE.match(requester_ident):
                    if requester_ident in taken_emails:
                        result.errors.append((row_no, f'Requester "{requester_ident}" belongs to another school.')); continue
                    username = make_unique_username(requester_ident.split('@')[0], school, reserved=taken_users)
                    creator = build_user(username=username, email=requester_ident, password=generate_password(),
                                         role=ROLE_USER, school=school, must_change_password=True)
                    db.session.add(creator)
                    db.session.flush()
                    taken_users.add(username.lower()); taken_emails.add(requester_ident)
                    cache['users'][username.lower()] = creator
                    cache['users'][requester_ident] = creator
                    result.new_users.append({'username': username, 'email': requester_ident, 'school': school.name})
                elif fallback_to_importer and (actor.school_id == school.id or actor.role == ROLE_SUPER_ADMIN):
                    creator = actor
                else:
                    shown = _s(row.get('requester')) or '(blank)'
                    result.errors.append((row_no, f'Requester {shown} not found in {school.name}. '
                                                  'Import the user first, or tick "create missing requesters".'))
                    continue

            # assignee
            assignee = None
            assignee_ident = _s(row.get('assignee')).lower()
            if assignee_ident:
                assignee = cache['users'].get(assignee_ident)
                if assignee is None or assignee.role not in AGENT_ROLES:
                    result.errors.append((row_no, f'Assignee "{_s(row.get("assignee"))}" is not an agent in {school.name}.')); continue
                if assignee.role != CATEGORY_AGENT_ROLE[category] and not (category == 'HSE' and assignee.role == 'fm_agent'):
                    result.errors.append((row_no, f'Assignee "{assignee.username}" cannot handle {category} tickets.')); continue
            if assignee and status == 'Open':
                status = 'Assigned'
            if not assignee and status == 'Assigned':
                status = 'Open'

            # duplicates (re-running the same file must be harmless)
            ref = _s(row.get('external_ref'))[:100]
            if ref and ref.lower() in cache['refs']:
                result.skipped.append((row_no, f'Ticket reference "{ref}" already imported.')); continue
            key = (title.lower(), creator.id, created_at.replace(microsecond=0))
            if key in cache['keys']:
                result.skipped.append((row_no, 'Duplicate of an existing ticket (same title, requester and date).')); continue

            ticket = Ticket(
                title=title, description=description, category=category, priority=priority,
                status=status, created_by=creator.id, assigned_to=assignee.id if assignee else None,
                created_at=created_at, updated_at=updated_at, school_id=school.id,
                external_ref=ref or None)
            db.session.add(ticket)
            comment_text = _s(row.get('comments'))
            if comment_text:
                author = assignee or creator
                db.session.flush()
                db.session.add(Comment(ticket_id=ticket.id, user_id=author.id, content=comment_text,
                                       created_at=updated_at, school_id=school.id))
            if ref:
                cache['refs'].add(ref.lower())
            cache['keys'].add(key)
            result.created.append({'row': row_no, 'title': title, 'category': category, 'status': status,
                                   'requester': creator.username, 'school': school.name})
        except Exception as exc:
            result.errors.append((row_no, f'Unexpected error: {exc}'))

    _finish(result, actor, default_school, 'bulk_import_tickets',
            f'Imported {len(result.created)} tickets ({len(result.errors)} errors, {len(result.skipped)} skipped)')
    return result


def _finish(result, actor, school, action, details):
    """Commit (or roll back for a dry run) and write one audit entry."""
    if result.dry_run or not (result.created or result.new_users):
        db.session.rollback()
        return
    try:
        db.session.add(UserActivityLog(user_id=actor.id, username=actor.username, action=action,
                                       details=details, school_id=school.id if school else actor.school_id))
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        result.errors.append((0, f'Database error - nothing was imported: {exc}'))
        result.created.clear()
        result.new_users.clear()


# --------------------------------------------------------------------------- templates
USER_TEMPLATE_HEADERS = ['email', 'username', 'full_name', 'role', 'department', 'password', 'school']
USER_TEMPLATE_ROWS = [
    ['jane.doe@school.edu', 'jane.doe', 'Jane Doe', 'user', 'Science', '', 'my-school'],
    ['ali.it@school.edu', '', 'Ali Hassan', 'it_agent', 'IT', '', 'my-school'],
    ['sara.fm@school.edu', '', 'Sara Khan', 'fm_agent', 'Facilities', '', 'my-school'],
    ['omar.hse@school.edu', '', 'Omar Said', 'hse_agent', 'HSE', '', 'my-school'],
]
TICKET_TEMPLATE_HEADERS = [
    'ticket_id', 'title', 'description', 'category', 'priority', 'status', 'requester',
    'assigned_to', 'created_at', 'updated_at', 'comments', 'school']
TICKET_TEMPLATE_ROWS = [
    ['OLD-1001', 'Projector not working', 'Room 12 projector shows no signal', 'IT', 'High', 'Closed',
     'jane.doe@school.edu', 'ali.it@school.edu', '2025-09-03 08:15', '2025-09-04 10:00',
     'Replaced HDMI cable', 'my-school'],
    ['OLD-1002', 'Leaking tap in staff room', 'Tap leaks continuously', 'FM', 'Medium', 'Open',
     'jane.doe@school.edu', '', '03/10/2025 09:30', '', '', 'my-school'],
]


def template_csv(kind):
    headers, rows = ((USER_TEMPLATE_HEADERS, USER_TEMPLATE_ROWS) if kind == 'users'
                     else (TICKET_TEMPLATE_HEADERS, TICKET_TEMPLATE_ROWS))
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return buf.getvalue()


def template_xlsx(kind):
    headers, rows = ((USER_TEMPLATE_HEADERS, USER_TEMPLATE_ROWS) if kind == 'users'
                     else (TICKET_TEMPLATE_HEADERS, TICKET_TEMPLATE_ROWS))
    wb = Workbook()
    ws = wb.active
    ws.title = kind.capitalize()
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 24
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
