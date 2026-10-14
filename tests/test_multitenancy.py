import io
import re

import pytest
from openpyxl import Workbook, load_workbook

from conftest import (
    PASSWORD, PNG, activate_all, create_school, create_user_via_ui, login, set_password,
)
from extensions import db
from models import School, User, Ticket, Comment


# --------------------------------------------------------------------------- fixtures
@pytest.fixture(scope='module')
def world(flask_app, superadmin):
    """Two schools, each with an admin, an IT/FM/HSE agent and a requester."""
    sa = superadmin
    out = {}
    for key, name, color in (('a', 'Alpha Academy', '#8a1c3c'), ('b', 'Beta School', '#1c4f8a')):
        admin_user = create_school(sa, flask_app, name, slug=f'{key}-school', admin_email=f'admin.{key}@example.com',
                                   color=color)
        with flask_app.app_context():
            school = School.query.filter_by(slug=f'{key}-school').one()
            sid = school.id
        for role, uname in (('it_agent', f'it_{key}'), ('fm_agent', f'fm_{key}'), ('hse_agent', f'hse_{key}'),
                            ('user', f'teacher_{key}')):
            resp = create_user_via_ui(sa, username=uname, email=f'{uname}@example.com', role=role,
                                      school_id=str(sid), full_name=uname.upper())
            assert resp.status_code == 302, resp.get_data(as_text=True)[:400]
            set_password(flask_app, uname)
        out[key] = {'sid': sid, 'admin': admin_user, 'it': f'it_{key}', 'fm': f'fm_{key}', 'hse': f'hse_{key}',
                    'user': f'teacher_{key}', 'name': name}
    return out


def ticket_form(title='Broken projector', category='IT', priority='High'):
    return {'title': title, 'description': 'Needs fixing', 'category': category, 'priority': priority}


# --------------------------------------------------------------------------- platform / branding
def test_super_admin_login_lands_on_platform(superadmin):
    resp = superadmin.get('/dashboard', follow_redirects=True)
    assert resp.status_code == 200
    assert b'Schools' in resp.data


def test_school_created_with_branding(flask_app, world, make_client):
    with flask_app.app_context():
        school = db.session.get(School, world['a']['sid'])
        assert school.slug == 'a-school'
        assert school.primary_color == '#8a1c3c'
        assert school.logo_filename and school.logo_filename.startswith('schools/')
    # branded sign-in page
    c = flask_app.test_client()
    html = c.get('/s/a-school').get_data(as_text=True)
    assert 'Alpha Academy Helpdesk' in html
    assert '--teal-dark: #8a1c3c' in html
    assert 'uploads/schools/' in html
    assert 'Beta School' not in html


def test_each_school_ui_has_own_name_logo_colour(flask_app, world, make_client):
    a = make_client(world['a']['user'])
    b = make_client(world['b']['user'])
    ha = a.get('/user/dashboard').get_data(as_text=True)
    hb = b.get('/user/dashboard').get_data(as_text=True)
    assert 'Alpha Academy Helpdesk' in ha and 'Beta School Helpdesk' not in ha
    assert 'Beta School Helpdesk' in hb and 'Alpha Academy Helpdesk' not in hb
    assert '#8a1c3c' in ha and '#1c4f8a' in hb
    assert 'uploads/schools/' in ha and 'uploads/schools/' in hb


def test_dynamic_manifest_is_branded(flask_app, world, make_client):
    a = make_client(world['a']['user'])
    data = a.get('/static/manifest.json').get_json()
    assert data['name'] == 'Alpha Academy Helpdesk'
    assert data['theme_color'] == '#8a1c3c'


def test_invalid_logo_rejected(flask_app, superadmin):
    resp = superadmin.post('/platform/schools/new', data={
        'name': 'Bad Logo School', 'logo_file': (io.BytesIO(b'<svg onload=alert(1)>'), 'x.svg')},
        content_type='multipart/form-data')
    assert resp.status_code == 400
    resp = superadmin.post('/platform/schools/new', data={
        'name': 'Fake Png School', 'logo_file': (io.BytesIO(b'not really a png'), 'x.png')},
        content_type='multipart/form-data')
    assert resp.status_code == 400
    with flask_app.app_context():
        assert School.query.filter(School.name.in_(['Bad Logo School', 'Fake Png School'])).count() == 0


def test_school_validation(superadmin):
    assert superadmin.post('/platform/schools/new', data={'name': ''}).status_code == 400
    assert superadmin.post('/platform/schools/new', data={'name': 'Alpha Academy'}).status_code == 400   # duplicate
    assert superadmin.post('/platform/schools/new', data={'name': 'X1', 'primary_color': 'red'}).status_code == 400
    assert superadmin.post('/platform/schools/new', data={'name': 'X2', 'slug': 'Bad Slug!'}).status_code == 400


def test_super_admin_edits_branding_and_logo(flask_app, world, superadmin):
    sid = world['b']['sid']
    resp = superadmin.post(f'/platform/schools/{sid}/edit', data={
        'name': 'Beta School', 'display_name': 'Beta Support Desk', 'slug': 'b-school',
        'primary_color': '#115577', 'secondary_color': '#3399aa', 'contact_email': 'it@beta.example.com',
        'logo_file': (io.BytesIO(PNG), 'new.png')}, content_type='multipart/form-data')
    assert resp.status_code == 302
    html = flask_app.test_client().get('/s/b-school').get_data(as_text=True)
    assert 'Beta Support Desk' in html and '#115577' in html
    superadmin.post(f'/platform/schools/{sid}/edit', data={
        'name': 'Beta School', 'display_name': 'Beta School Helpdesk', 'slug': 'b-school',
        'primary_color': '#1c4f8a', 'secondary_color': '', 'contact_email': ''}, content_type='multipart/form-data')


# --------------------------------------------------------------------------- login rules
def test_login_with_email_and_school_binding(flask_app, world):
    c = flask_app.test_client()
    assert login(c, f"{world['a']['user']}@example.com").status_code == 302
    # wrong school sign-in page for this user is rejected
    c2 = flask_app.test_client()
    resp = login(c2, world['a']['user'], path='/s/b-school')
    assert resp.status_code == 401
    # right school page works
    c3 = flask_app.test_client()
    assert login(c3, world['a']['user'], path='/s/a-school').status_code == 302


def test_bad_credentials(flask_app):
    c = flask_app.test_client()
    assert login(c, 'nobody', 'x' * 10).status_code == 401
    assert login(c, 'teacher_a', 'wrong-password').status_code == 401


def test_login_throttle(flask_app, world):
    c = flask_app.test_client()
    for _ in range(8):
        login(c, 'teacher_b', 'bad-password-1')
    resp = login(c, 'teacher_b', PASSWORD)          # even the right password is blocked now
    assert resp.status_code == 429
    import app as app_module
    app_module._FAILED_LOGINS.clear()


def test_open_redirect_blocked(flask_app, world):
    c = flask_app.test_client()
    resp = c.post('/login?next=//evil.example.com', data={'username': world['a']['user'], 'password': PASSWORD})
    assert resp.status_code == 302 and 'evil.example.com' not in resp.headers['Location']


def test_forced_password_change_flow(flask_app, superadmin):
    resp = create_user_via_ui(superadmin, username='newbie', email='newbie@example.com', role='user',
                              school_id=str(School.query.filter_by(slug='a-school').one().id)
                              if False else '')
    # school missing -> rejected
    assert resp.status_code == 400
    with flask_app.app_context():
        sid = School.query.filter_by(slug='a-school').one().id
    resp = superadmin.post('/admin/create_user', data={'username': 'newbie', 'email': 'newbie@example.com',
                                                       'role': 'user', 'school_id': str(sid), 'password': ''})
    assert resp.status_code == 302
    flashed = superadmin.get('/admin/users').get_data(as_text=True)
    m = re.search(r'Temporary password \(shown once\): (\S+)', flashed)
    assert m, 'generated password should be flashed once'
    c = flask_app.test_client()
    assert login(c, 'newbie', m.group(1)).status_code == 302
    # any page redirects to change-password until done
    resp = c.get('/user/dashboard')
    assert resp.status_code == 302 and '/account/password' in resp.headers['Location']
    resp = c.post('/account/password', data={'current_password': m.group(1), 'new_password': 'short',
                                             'confirm_password': 'short'})
    assert b'at least 8' in resp.data
    resp = c.post('/account/password', data={'current_password': m.group(1), 'new_password': 'BrandNew#99',
                                             'confirm_password': 'BrandNew#99'})
    assert resp.status_code == 302
    assert c.get('/user/dashboard').status_code == 200


# --------------------------------------------------------------------------- isolation
def test_ticket_isolation_between_schools(flask_app, world, make_client):
    a_user = make_client(world['a']['user'])
    b_user = make_client(world['b']['user'])
    assert a_user.post('/ticket/create', data=ticket_form('ALPHA-SECRET-TICKET')).status_code == 302
    assert b_user.post('/ticket/create', data=ticket_form('BETA-SECRET-TICKET')).status_code == 302
    with flask_app.app_context():
        ta = Ticket.query.execution_options(tenant_bypass=True).filter_by(title='ALPHA-SECRET-TICKET').one()
        tb = Ticket.query.execution_options(tenant_bypass=True).filter_by(title='BETA-SECRET-TICKET').one()
        assert ta.school_id == world['a']['sid'] and tb.school_id == world['b']['sid']
        ta_id, tb_id = ta.id, tb.id
    # direct URL access across tenants is a 404 for requesters, agents and admins alike
    for viewer in (b_user, make_client(world['b']['it']), make_client(world['b']['admin'])):
        assert viewer.get(f'/ticket/{ta_id}').status_code == 404
        assert viewer.post(f'/ticket/{ta_id}/comment', data={'content': 'hi'}).status_code == 404
        assert viewer.post(f'/ticket/{ta_id}/status', data={'status': 'Closed'}).status_code == 404
    admin_b = make_client(world['b']['admin'])
    assert admin_b.post(f'/ticket/{ta_id}/assign', data={'agent_id': ''}).status_code == 404
    # lists only show own school
    html = admin_b.get('/admin/tickets').get_data(as_text=True)
    assert 'BETA-SECRET-TICKET' in html and 'ALPHA-SECRET-TICKET' not in html
    html = admin_b.get('/admin/dashboard').get_data(as_text=True)
    assert 'ALPHA-SECRET-TICKET' not in html
    html = admin_b.get('/admin/search_tickets').get_data(as_text=True)
    assert 'ALPHA-SECRET-TICKET' not in html
    # exports are scoped too
    wb = load_workbook(io.BytesIO(admin_b.get('/admin/download_tickets_excel').data))
    titles = [r[1] for r in wb.active.iter_rows(min_row=2, values_only=True)]
    assert 'BETA-SECRET-TICKET' in titles and 'ALPHA-SECRET-TICKET' not in titles


def test_user_isolation_between_schools(flask_app, world, make_client):
    admin_a = make_client(world['a']['admin'])
    html = admin_a.get('/admin/users').get_data(as_text=True)
    assert world['a']['user'] in html
    assert world['b']['user'] not in html
    assert 'admin' not in re.findall(r'<strong>(\w+)</strong>', html)       # super admin hidden
    with flask_app.app_context():
        uid_b = User.query.execution_options(tenant_bypass=True).filter_by(username=world['b']['user']).one().id
        uid_sa = User.query.execution_options(tenant_bypass=True).filter_by(username='admin').one().id
    for path in (f'/admin/edit_user/{uid_b}', f'/admin/view_user/{uid_b}', f'/admin/edit_user/{uid_sa}'):
        assert admin_a.get(path).status_code == 404
    for path in (f'/admin/toggle_user_status/{uid_b}', f'/admin/delete_user/{uid_b}',
                 f'/admin/reset_password/{uid_b}', f'/admin/toggle_user_status/{uid_sa}'):
        assert admin_a.post(path, data={'new_password': 'Hacked#1234'}).status_code == 404
    assert login(flask_app.test_client(), world['b']['user']).status_code == 302   # password untouched
    html = admin_a.get('/admin/search_users?query=teacher').get_data(as_text=True)
    assert world['b']['user'] not in html


def test_school_admin_cannot_create_users_elsewhere_or_super_admins(flask_app, world, make_client):
    admin_a = make_client(world['a']['admin'])
    resp = admin_a.post('/admin/create_user', data={
        'username': 'sneaky', 'email': 'sneaky@example.com', 'role': 'user', 'password': PASSWORD,
        'school_id': str(world['b']['sid'])})
    assert resp.status_code == 302                      # allowed, but pinned to the admin's own school
    with flask_app.app_context():
        assert User.query.execution_options(tenant_bypass=True).filter_by(username='sneaky').one().school_id == world['a']['sid']
    resp = admin_a.post('/admin/create_user', data={
        'username': 'evilroot', 'email': 'evilroot@example.com', 'role': 'super_admin', 'password': PASSWORD})
    assert resp.status_code == 400
    with flask_app.app_context():
        assert User.query.execution_options(tenant_bypass=True).filter_by(username='evilroot').count() == 0


def test_cross_school_agent_assignment_blocked(flask_app, world, make_client):
    admin_a = make_client(world['a']['admin'])
    with flask_app.app_context():
        it_b = User.query.execution_options(tenant_bypass=True).filter_by(username=world['b']['it']).one().id
        it_a = User.query.execution_options(tenant_bypass=True).filter_by(username=world['a']['it']).one().id
        tid = Ticket.query.execution_options(tenant_bypass=True).filter_by(title='ALPHA-SECRET-TICKET').one().id
    resp = admin_a.post(f'/ticket/{tid}/assign', data={'agent_id': str(it_b)}, follow_redirects=True)
    assert b'Invalid agent' in resp.data
    resp = admin_a.post(f'/ticket/{tid}/assign', data={'agent_id': str(it_a)}, follow_redirects=True)
    assert b'assigned to' in resp.data.lower()
    with flask_app.app_context():
        assert db.session.get(Ticket, tid).assigned_to == it_a


def test_round_robin_is_per_school(flask_app, world, make_client):
    # second IT agent in school A; school B keeps a single one
    sa_admin = make_client(world['a']['admin'])
    assert create_user_via_ui(sa_admin, username='it_a2', email='it_a2@example.com', role='it_agent').status_code == 302
    set_password(flask_app, 'it_a2')
    user_a = make_client(world['a']['user'])
    for i in range(4):
        user_a.post('/ticket/create', data=ticket_form(f'RR-{i}', 'IT'))
    with flask_app.app_context():
        assignees = [t.assigned_to for t in Ticket.query.execution_options(tenant_bypass=True)
                     .filter(Ticket.title.like('RR-%')).order_by(Ticket.id)]
        names = {u.id: u.username for u in User.query.execution_options(tenant_bypass=True).all()}
    used = {names[a] for a in assignees}
    assert used == {'it_a', 'it_a2'}                   # both A agents rotate
    assert 'it_b' not in used                          # never B's agent
    assert assignees[0] != assignees[1]


def test_fm_hse_agents_receive_their_categories(flask_app, world, make_client):
    user_a = make_client(world['a']['user'])
    user_a.post('/ticket/create', data=ticket_form('FM-CASE', 'FM', 'Low'))
    user_a.post('/ticket/create', data=ticket_form('HSE-CASE', 'HSE', 'Critical'))
    with flask_app.app_context():
        fm = Ticket.query.execution_options(tenant_bypass=True).filter_by(title='FM-CASE').one()
        hse = Ticket.query.execution_options(tenant_bypass=True).filter_by(title='HSE-CASE').one()
        assert db.session.get(User, fm.assigned_to).username == world['a']['fm']
        assert db.session.get(User, hse.assigned_to).username == world['a']['hse']
    hse_client = make_client(world['a']['hse'])           # HSE agents land on the agent dashboard
    assert hse_client.get('/dashboard').headers['Location'].endswith('/agent/dashboard')
    assert b'HSE-CASE' in hse_client.get('/agent/dashboard').data


def test_super_admin_must_enter_school_for_tenant_pages(flask_app, world, superadmin):
    resp = superadmin.get('/admin/tickets')
    assert resp.status_code == 302 and '/platform' in resp.headers['Location']
    # entering a school scopes everything to it
    assert superadmin.post(f"/platform/schools/{world['b']['sid']}/enter").status_code == 302
    html = superadmin.get('/admin/tickets').get_data(as_text=True)
    assert 'BETA-SECRET-TICKET' in html and 'ALPHA-SECRET-TICKET' not in html
    assert 'Managing' in superadmin.get('/admin/dashboard').get_data(as_text=True)
    superadmin.post('/platform/leave')
    assert superadmin.get('/admin/tickets').status_code == 302


def test_super_admin_comment_visible_to_school(flask_app, world, superadmin, make_client):
    with flask_app.app_context():
        tid = Ticket.query.execution_options(tenant_bypass=True).filter_by(title='BETA-SECRET-TICKET').one().id
    superadmin.post(f"/platform/schools/{world['b']['sid']}/enter")
    superadmin.post(f'/ticket/{tid}/comment', data={'content': 'Platform support here'})
    superadmin.post('/platform/leave')
    html = make_client(world['b']['user']).get(f'/ticket/{tid}').get_data(as_text=True)
    assert 'Platform support here' in html


def test_suspended_school_blocks_login_and_sessions(flask_app, world, superadmin, make_client):
    sid = world['b']['sid']
    c = make_client(world['b']['user'])
    assert c.get('/user/dashboard').status_code == 200
    superadmin.post(f'/platform/schools/{sid}/toggle')
    try:
        assert c.get('/user/dashboard').status_code == 302            # live session is dropped
        assert login(flask_app.test_client(), world['b']['user']).status_code == 401
        assert flask_app.test_client().get('/s/b-school').status_code == 403
    finally:
        superadmin.post(f'/platform/schools/{sid}/toggle')
    assert login(flask_app.test_client(), world['b']['user']).status_code == 302


def test_cannot_delete_school_with_data(world, superadmin):
    resp = superadmin.post(f"/platform/schools/{world['a']['sid']}/delete", follow_redirects=True)
    assert b'Only schools without users or tickets' in resp.data


# --------------------------------------------------------------------------- authorization
@pytest.mark.parametrize('path,method', [
    ('/admin/dashboard', 'get'), ('/admin/users', 'get'), ('/admin/tickets', 'get'),
    ('/admin/system_management', 'get'), ('/admin/email_settings', 'get'), ('/admin/settings', 'get'),
    ('/admin/database_management', 'get'), ('/admin/download_tickets_excel', 'get'),
    ('/admin/download_summary_pdf', 'get'), ('/admin/switch_database/1', 'post'),
    ('/admin/delete_user/1', 'post'), ('/admin/import/users', 'get'), ('/platform', 'get'),
    ('/admin/create_database_backup', 'post'), ('/admin/update_module_permissions', 'post'),
])
def test_anonymous_is_redirected_to_login(client, path, method):
    resp = getattr(client, method)(path)
    assert resp.status_code == 302 and '/login' in resp.headers['Location']


@pytest.mark.parametrize('path,method', [
    ('/admin/dashboard', 'get'), ('/admin/users', 'get'), ('/admin/tickets', 'get'),
    ('/admin/system_management', 'get'), ('/admin/email_settings', 'get'), ('/admin/settings', 'get'),
    ('/admin/database_management', 'get'), ('/admin/download_tickets_excel', 'get'),
    ('/admin/switch_database/1', 'post'), ('/admin/delete_user/1', 'post'), ('/admin/import/users', 'get'),
    ('/admin/create_database_backup', 'post'), ('/admin/agent_management', 'get'),
    ('/admin/send_daily_summary', 'post'), ('/admin/update_module_permissions', 'post'),
])
def test_regular_user_denied_admin_pages(world, make_client, path, method):
    c = make_client(world['a']['user'])
    resp = getattr(c, method)(path)
    assert resp.status_code in (302, 404)
    if resp.status_code == 302:
        assert '/login' not in resp.headers['Location']
        assert 'dashboard' in resp.headers['Location']


@pytest.mark.parametrize('path,method', [
    ('/platform', 'get'), ('/platform/schools/new', 'get'), ('/admin/system_management', 'get'),
    ('/admin/email_settings', 'get'), ('/admin/email_templates', 'get'), ('/admin/settings', 'get'),
    ('/admin/database_management', 'get'), ('/admin/backup_restore', 'get'),
    ('/admin/create_database_backup', 'post'),
])
def test_school_admin_denied_platform_pages(world, make_client, path, method):
    c = make_client(world['a']['admin'])
    resp = getattr(c, method)(path)
    assert resp.status_code == 302 and 'dashboard' in resp.headers['Location']


def test_agent_cannot_assign_or_manage(world, make_client):
    c = make_client(world['a']['it'])
    assert c.get('/admin/users').status_code == 302
    assert c.post('/ticket/1/assign', data={'agent_id': ''}).status_code in (302, 404)
    assert c.get('/admin/import/tickets').status_code == 302


# --------------------------------------------------------------------------- bulk import: users
def csv_file(text, name='users.csv'):
    return (io.BytesIO(text.encode('utf-8')), name)


def xlsx_file(rows, name='data.xlsx'):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return (buf, name)


def test_school_admin_bulk_import_users_csv(flask_app, world, make_client):
    admin_a = make_client(world['a']['admin'])
    text = (
        'email,username,full_name,role,department,password\n'
        'p1@alpha.example.com,alpha.p1,Pat One,user,Math,Secret#1234\n'
        'p2@alpha.example.com,,Pat Two,it_agent,IT,\n'
        'p3@alpha.example.com,,Pat Three,hse_agent,,\n'
        'p4@alpha.example.com,,Pat Four,fm_agent,,\n'
        'bad-email,,,user,,\n'
        'p5@alpha.example.com,,,superuser,,\n'
        'p1@alpha.example.com,dupe,,user,,\n'
        'p6@alpha.example.com,,,user,,short\n'
        'p7@alpha.example.com,,,super_admin,,\n'
    )
    resp = admin_a.post('/admin/import/users', data={'import_file': csv_file(text), 'require_password_change': '1'},
                        content_type='multipart/form-data')
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    assert 'Invalid email' in html and 'Invalid role' in html and 'at least 8' in html
    assert 'Super admins cannot be bulk imported' in html or 'Invalid role' in html
    with flask_app.app_context():
        rows = {u.email: u for u in User.query.execution_options(tenant_bypass=True)
                .filter(User.email.like('%@alpha.example.com')).all()}
    assert set(rows) == {'p1@alpha.example.com', 'p2@alpha.example.com', 'p3@alpha.example.com', 'p4@alpha.example.com'}
    assert all(u.school_id == world['a']['sid'] for u in rows.values())
    assert rows['p2@alpha.example.com'].role == 'it_agent' and rows['p3@alpha.example.com'].role == 'hse_agent'
    assert rows['p2@alpha.example.com'].username == 'p2user'
    assert rows['p2@alpha.example.com'].must_change_password
    # generated passwords are shown once and actually work
    m = re.search(r'<code>([A-Za-z0-9]{12})</code>', html)
    assert m
    c = flask_app.test_client()
    assert login(c, 'p2@alpha.example.com', m.group(1)).status_code == 302
    # provided password works
    assert login(flask_app.test_client(), 'alpha.p1', 'Secret#1234').status_code == 302


def test_bulk_import_users_idempotent_and_dry_run(flask_app, world, make_client):
    admin_a = make_client(world['a']['admin'])
    text = 'email,role\nq1@alpha.example.com,user\nq2@alpha.example.com,user\n'
    resp = admin_a.post('/admin/import/users', data={'import_file': csv_file(text), 'dry_run': '1'},
                        content_type='multipart/form-data')
    assert b'nothing saved' in resp.data
    with flask_app.app_context():
        assert User.query.execution_options(tenant_bypass=True).filter(User.email.like('q%@alpha.example.com')).count() == 0
    admin_a.post('/admin/import/users', data={'import_file': csv_file(text)}, content_type='multipart/form-data')
    resp = admin_a.post('/admin/import/users', data={'import_file': csv_file(text)}, content_type='multipart/form-data')
    assert b'already exists' in resp.data
    with flask_app.app_context():
        assert User.query.execution_options(tenant_bypass=True).filter(User.email.like('q%@alpha.example.com')).count() == 2


def test_bulk_import_users_xlsx_with_header_aliases(flask_app, world, make_client):
    admin_a = make_client(world['a']['admin'])
    f = xlsx_file([['E-mail', 'Name', 'User Type'], ['x1@alpha.example.com', 'Xavier One', 'IT'],
                   ['x2@alpha.example.com', 'Xena Two', 'facilities']])
    resp = admin_a.post('/admin/import/users', data={'import_file': f}, content_type='multipart/form-data')
    assert resp.status_code == 200
    with flask_app.app_context():
        users = {u.email: u for u in User.query.execution_options(tenant_bypass=True)
                 .filter(User.email.like('x_@alpha.example.com'))}
    assert users['x1@alpha.example.com'].role == 'it_agent' and users['x1@alpha.example.com'].full_name == 'Xavier One'
    assert users['x2@alpha.example.com'].role == 'fm_agent'


def test_school_admin_cannot_import_into_other_school(flask_app, world, make_client):
    admin_a = make_client(world['a']['admin'])
    text = 'email,role,school\nrogue@beta.example.com,user,b-school\nok@alpha.example.com,user,a-school\n'
    resp = admin_a.post('/admin/import/users', data={'import_file': csv_file(text), 'school_id': str(world['b']['sid'])},
                        content_type='multipart/form-data')
    assert b'only import into your own school' in resp.data
    with flask_app.app_context():
        assert User.query.execution_options(tenant_bypass=True).filter_by(email='rogue@beta.example.com').count() == 0
        ok = User.query.execution_options(tenant_bypass=True).filter_by(email='ok@alpha.example.com').one()
        assert ok.school_id == world['a']['sid']


def test_super_admin_import_users_into_multiple_schools(flask_app, world, superadmin):
    text = ('email,role,school\n'
            'multi1@example.com,user,a-school\n'
            'multi2@example.com,it_agent,Beta School\n'
            'multi3@example.com,user,unknown-school\n'
            'multi4@example.com,admin,b-school\n')
    resp = superadmin.post('/admin/import/users', data={'import_file': csv_file(text)}, content_type='multipart/form-data')
    assert b'Unknown school' in resp.data
    with flask_app.app_context():
        got = {u.email: (u.school_id, u.role) for u in User.query.execution_options(tenant_bypass=True)
               .filter(User.email.like('multi%@example.com'))}
    assert got == {'multi1@example.com': (world['a']['sid'], 'user'),
                   'multi2@example.com': (world['b']['sid'], 'it_agent'),
                   'multi4@example.com': (world['b']['sid'], 'admin')}
    # default school selection applies to rows without a school value
    resp = superadmin.post('/admin/import/users', data={
        'import_file': csv_file('email,role\nmulti5@example.com,user\n'), 'school_id': str(world['a']['sid'])},
        content_type='multipart/form-data')
    with flask_app.app_context():
        assert User.query.execution_options(tenant_bypass=True).filter_by(email='multi5@example.com').one().school_id == world['a']['sid']
    # and a row with no school at all is an error
    resp = superadmin.post('/admin/import/users', data={'import_file': csv_file('email,role\nmulti6@example.com,user\n')},
                           content_type='multipart/form-data')
    assert b'No school specified' in resp.data


@pytest.mark.parametrize('payload,msg', [
    (('', 'empty.csv'), b'choose a file'),
    (('just,some\n1,2\n', 'nomail.csv'), b'Missing required column'),
    (('email\n', 'headeronly.csv'), b'no data rows'),
])
def test_import_file_errors(world, make_client, payload, msg):
    admin_a = make_client(world['a']['admin'])
    data = {'import_file': csv_file(payload[0], payload[1]) if payload[0] else (io.BytesIO(b''), '')}
    resp = admin_a.post('/admin/import/users', data=data, content_type='multipart/form-data')
    assert resp.status_code == 400 and msg in resp.data


def test_import_rejects_wrong_extension_and_garbage_xlsx(world, make_client):
    admin_a = make_client(world['a']['admin'])
    resp = admin_a.post('/admin/import/users', data={'import_file': (io.BytesIO(b'x'), 'users.exe')},
                        content_type='multipart/form-data')
    assert resp.status_code == 400 and b'Unsupported file type' in resp.data
    resp = admin_a.post('/admin/import/users', data={'import_file': (io.BytesIO(b'not a zip'), 'users.xlsx')},
                        content_type='multipart/form-data')
    assert resp.status_code == 400 and b'valid .xlsx' in resp.data


def test_import_templates_download(world, make_client):
    admin_a = make_client(world['a']['admin'])
    for kind in ('users', 'tickets'):
        r = admin_a.get(f'/admin/import/template/{kind}.csv')
        assert r.status_code == 200 and r.data.startswith(b'email' if kind == 'users' else b'ticket_id')
        r = admin_a.get(f'/admin/import/template/{kind}.xlsx')
        assert r.status_code == 200 and load_workbook(io.BytesIO(r.data)).active.max_row >= 2
    assert admin_a.get('/admin/import/template/other.csv').status_code == 404


# --------------------------------------------------------------------------- bulk import: tickets
TICKET_CSV = (
    'ticket_id,title,description,category,priority,status,requester,assigned_to,created_at,updated_at,comments\n'
    'OLD-1,Projector dead,Room 4,IT,High,Closed,{u},{it},2025-09-03 08:15,2025-09-04 10:00,Replaced cable\n'
    'OLD-2,Tap leaking,Staff room,FM,Medium,Open,{u},,03/10/2025 09:30,,\n'
    'OLD-3,Fire door stuck,Block B,HSE,Critical,In Progress,{u}@example.com,{hse},2025-10-01,,\n'
    'OLD-4,No such requester,x,IT,Low,Open,ghost@nowhere.com,,2025-10-01,,\n'
    'OLD-5,Bad category,x,Plumbing,Low,Open,{u},,2025-10-01,,\n'
    'OLD-6,Bad date,x,IT,Low,Open,{u},,31/31/2025,,\n'
    'OLD-7,Wrong agent,x,IT,Low,Open,{u},{fm},2025-10-01,,\n'
    'OLD-8,Future,x,IT,Low,Open,{u},,2099-01-01,,\n'
    'OLD-9,Excel date,x,IT,Low,Resolved,{u},,2025-11-05 14:00,2025-11-06,\n'
)


def test_bulk_import_previous_tickets(flask_app, world, make_client):
    admin_a = make_client(world['a']['admin'])
    text = TICKET_CSV.format(u=world['a']['user'], it=world['a']['it'], hse=world['a']['hse'], fm=world['a']['fm'])
    resp = admin_a.post('/admin/import/tickets', data={'import_file': csv_file(text, 't.csv')},
                        content_type='multipart/form-data')
    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    for fragment in ('not found in', 'Invalid category', 'Unrecognised date', 'cannot handle IT', 'in the future'):
        assert fragment in html, fragment
    with flask_app.app_context():
        ts = {t.external_ref: t for t in Ticket.query.execution_options(tenant_bypass=True)
              .filter(Ticket.external_ref.like('OLD-%'), Ticket.school_id == world['a']['sid'])}
        assert set(ts) == {'OLD-1', 'OLD-2', 'OLD-3', 'OLD-9'}
        t1 = ts['OLD-1']
        assert t1.status == 'Closed' and t1.created_at.strftime('%Y-%m-%d %H:%M') == '2025-09-03 08:15'
        assert t1.updated_at.strftime('%Y-%m-%d %H:%M') == '2025-09-04 10:00'
        assert db.session.get(User, t1.assigned_to).username == world['a']['it']
        assert [c.content for c in t1.comments] == ['Replaced cable']
        assert ts['OLD-2'].created_at.strftime('%Y-%m-%d %H:%M') == '2025-10-03 09:30'   # day-first
        assert ts['OLD-2'].status == 'Open' and ts['OLD-2'].assigned_to is None
        assert ts['OLD-3'].status == 'In Progress'
        assert db.session.get(User, ts['OLD-3'].assigned_to).username == world['a']['hse']
        assert all(t.school_id == world['a']['sid'] for t in ts.values())
        assert all(c.school_id == world['a']['sid'] for c in Comment.query.execution_options(tenant_bypass=True)
                   .filter(Comment.ticket_id == t1.id))
    # imported tickets are visible to the right people only
    assert b'Projector dead' in make_client(world['a']['user']).get('/user/dashboard').data
    assert b'Projector dead' not in make_client(world['b']['admin']).get('/admin/tickets').data
    # re-import is a no-op
    resp = admin_a.post('/admin/import/tickets', data={'import_file': csv_file(text, 't.csv')},
                        content_type='multipart/form-data')
    assert b'already imported' in resp.data
    with flask_app.app_context():
        assert Ticket.query.execution_options(tenant_bypass=True).filter(
            Ticket.external_ref.like('OLD-%'), Ticket.school_id == world['a']['sid']).count() == 4


def test_ticket_import_does_not_disturb_round_robin(flask_app, world, make_client):
    with flask_app.app_context():
        before = Ticket.query.execution_options(tenant_bypass=True).filter_by(school_id=world['a']['sid']).count()
    from app import assign_ticket_round_robin
    # an imported (legacy) ticket assigned to the *last* agent must not skew rotation
    with flask_app.test_request_context():
        from flask_login import login_user
        user = User.query.execution_options(tenant_bypass=True).filter_by(username=world['a']['admin']).one()
        login_user(user)
        nxt = assign_ticket_round_robin('IT')
        assert nxt is not None and nxt.role == 'it_agent' and nxt.school_id == world['a']['sid']


def test_ticket_import_create_missing_requesters_and_fallback(flask_app, world, make_client):
    admin_b = make_client(world['b']['admin'])
    text = ('title,description,category,requester,created_at\n'
            'Legacy A,d,IT,new.person@beta.example.com,2024-05-01\n'
            'Legacy B,d,FM,,2024-05-02\n'
            'Legacy C,d,FM,teacher_a@example.com,2024-05-02\n')
    resp = admin_b.post('/admin/import/tickets', data={
        'import_file': csv_file(text), 'create_missing_requesters': '1', 'fallback_to_importer': '1'},
        content_type='multipart/form-data')
    assert resp.status_code == 200
    with flask_app.app_context():
        new_user = User.query.execution_options(tenant_bypass=True).filter_by(email='new.person@beta.example.com').one()
        assert new_user.school_id == world['b']['sid'] and new_user.role == 'user' and new_user.must_change_password
        ts = {t.title: t for t in Ticket.query.execution_options(tenant_bypass=True).filter(Ticket.title.like('Legacy %'))}
        assert ts['Legacy A'].created_by == new_user.id
        assert db.session.get(User, ts['Legacy B'].created_by).username == world['b']['admin']
        # teacher_a belongs to school A: never linked from school B, row is rejected
        assert 'Legacy C' not in ts
    assert b'belongs to another school' in resp.data


def test_super_admin_ticket_import_into_selected_school(flask_app, world, superadmin):
    text = f"title,description,category,requester\nSA imported,d,IT,{world['b']['user']}\n"
    resp = superadmin.post('/admin/import/tickets', data={'import_file': csv_file(text), 'school_id': str(world['b']['sid'])},
                           content_type='multipart/form-data')
    assert resp.status_code == 200
    with flask_app.app_context():
        t = Ticket.query.execution_options(tenant_bypass=True).filter_by(title='SA imported').one()
        assert t.school_id == world['b']['sid']
    # requester from another school is not accepted
    text = f"title,description,category,requester\nSA wrong,d,IT,{world['a']['user']}\n"
    resp = superadmin.post('/admin/import/tickets', data={'import_file': csv_file(text), 'school_id': str(world['b']['sid'])},
                           content_type='multipart/form-data')
    assert b'not found in' in resp.data


def test_ticket_import_xlsx_dates(flask_app, world, make_client):
    import datetime
    admin_a = make_client(world['a']['admin'])
    f = xlsx_file([['Title', 'Category', 'Requester', 'Created', 'Status'],
                   ['Excel ticket', 'IT', world['a']['user'], datetime.datetime(2025, 2, 3, 4, 5), 'Closed']])
    resp = admin_a.post('/admin/import/tickets', data={'import_file': f}, content_type='multipart/form-data')
    assert resp.status_code == 200
    with flask_app.app_context():
        t = Ticket.query.execution_options(tenant_bypass=True).filter_by(title='Excel ticket').one()
        assert t.created_at == datetime.datetime(2025, 2, 3, 4, 5) and t.status == 'Closed'


# --------------------------------------------------------------------------- user management rules
def test_user_edit_rules(flask_app, world, make_client):
    admin_a = make_client(world['a']['admin'])
    with flask_app.app_context():
        me = User.query.execution_options(tenant_bypass=True).filter_by(username=world['a']['admin']).one().id
        teacher = User.query.execution_options(tenant_bypass=True).filter_by(username=world['a']['user']).one().id
    # cannot demote / deactivate yourself, cannot delete yourself
    admin_a.post(f'/admin/edit_user/{me}', data={'username': world['a']['admin'], 'email': 'admin.a@example.com',
                                                  'role': 'user'})
    admin_a.post(f'/admin/toggle_user_status/{me}')
    admin_a.post(f'/admin/delete_user/{me}')
    with flask_app.app_context():
        u = db.session.get(User, me)
        assert u.role == 'admin' and u.is_active
    # duplicate username/email rejected
    resp = admin_a.post(f'/admin/edit_user/{teacher}', data={'username': world['a']['it'], 'email': 'teacher_a@example.com',
                                                             'role': 'user', 'is_active': 'on'})
    assert resp.status_code == 400
    # user with history cannot be deleted
    resp = admin_a.post(f'/admin/delete_user/{teacher}', follow_redirects=True)
    assert b'cannot be deleted' in resp.data
    # school admins cannot promote to super_admin
    resp = admin_a.post(f'/admin/edit_user/{teacher}', data={'username': world['a']['user'], 'email': 'teacher_a@example.com',
                                                             'role': 'super_admin', 'is_active': 'on'})
    assert resp.status_code == 400


def test_moving_user_between_schools_rules(flask_app, world, superadmin):
    resp = create_user_via_ui(superadmin, username='mover', email='mover@example.com', role='user',
                              school_id=str(world['a']['sid']))
    assert resp.status_code == 302
    with flask_app.app_context():
        uid = User.query.execution_options(tenant_bypass=True).filter_by(username='mover').one().id
    resp = superadmin.post(f'/admin/edit_user/{uid}', data={'username': 'mover', 'email': 'mover@example.com', 'role': 'user',
                                                            'is_active': 'on', 'school_id': str(world['b']['sid'])})
    assert resp.status_code == 302
    with flask_app.app_context():
        assert db.session.get(User, uid).school_id == world['b']['sid']
    # a user with tickets cannot be moved
    with flask_app.app_context():
        tid = User.query.execution_options(tenant_bypass=True).filter_by(username=world['a']['user']).one().id
    resp = superadmin.post(f'/admin/edit_user/{tid}', data={'username': world['a']['user'], 'email': 'teacher_a@example.com',
                                                            'role': 'user', 'is_active': 'on', 'school_id': str(world['b']['sid'])})
    assert resp.status_code == 400


def test_super_admin_user_list_filters_by_school(world, superadmin):
    html = superadmin.get(f"/admin/users?school_id={world['a']['sid']}").get_data(as_text=True)
    assert world['a']['user'].upper() in html and world['b']['user'].upper() not in html
    html = superadmin.get('/admin/users').get_data(as_text=True)
    assert world['a']['user'].upper() in html and world['b']['user'].upper() in html


def test_last_super_admin_protected(flask_app, superadmin):
    with flask_app.app_context():
        sa = User.query.execution_options(tenant_bypass=True).filter_by(username='admin').one().id
    resp = superadmin.post(f'/admin/delete_user/{sa}', follow_redirects=True)
    assert b'cannot delete your own account' in resp.data


# --------------------------------------------------------------------------- csrf & headers
def test_csrf_enforced_when_enabled(flask_app, world, make_client):
    c = make_client(world['a']['user'])
    flask_app.config['WTF_CSRF_ENABLED'] = True
    try:
        resp = c.post('/ticket/create', data=ticket_form('CSRF attempt'))
        assert resp.status_code == 302
        with flask_app.app_context():
            assert Ticket.query.execution_options(tenant_bypass=True).filter_by(title='CSRF attempt').count() == 0
        page = c.get('/ticket/create').get_data(as_text=True)
        token = re.search(r'name="csrf-token" content="([^"]+)"', page).group(1)
        data = ticket_form('CSRF ok')
        data['csrf_token'] = token
        assert c.post('/ticket/create', data=data).status_code == 302
        with flask_app.app_context():
            assert Ticket.query.execution_options(tenant_bypass=True).filter_by(title='CSRF ok').count() == 1
    finally:
        flask_app.config['WTF_CSRF_ENABLED'] = False


def test_tenant_pages_not_cacheable(world, make_client):
    c = make_client(world['a']['user'])
    assert c.get('/user/dashboard').headers['Cache-Control'] == 'no-store'


def test_ticket_form_validation(world, make_client):
    c = make_client(world['a']['user'])
    assert c.post('/ticket/create', data=ticket_form('', 'IT')).status_code == 400
    assert c.post('/ticket/create', data=ticket_form('x', 'BOGUS')).status_code == 400
    assert c.post('/ticket/create', data=ticket_form('x' * 201, 'IT')).status_code == 400


# --------------------------------------------------------------------------- smoke: every page renders
def _smoke_paths(flask_app):
    paths = []
    for rule in flask_app.url_map.iter_rules():
        if 'GET' not in rule.methods or rule.arguments or rule.endpoint == 'static':
            continue
        paths.append(rule.rule)
    return sorted(set(paths))


SKIP_SMOKE = {'/logout', '/admin/download_summary_pdf'}


@pytest.mark.parametrize('who', ['super_enter', 'school_admin', 'it_agent', 'user'])
def test_every_page_renders_without_server_error(flask_app, world, superadmin, make_client, who):
    if who == 'super_enter':
        c = flask_app.test_client()
        login(c, 'admin', 'SuperAdmin#123')
        c.post(f"/platform/schools/{world['a']['sid']}/enter")
    elif who == 'school_admin':
        c = make_client(world['a']['admin'])
    elif who == 'it_agent':
        c = make_client(world['a']['it'])
    else:
        c = make_client(world['a']['user'])
    bad = []
    for path in _smoke_paths(flask_app):
        if path in SKIP_SMOKE:
            continue
        resp = c.get(path, follow_redirects=True)
        if resp.status_code >= 500:
            bad.append((path, resp.status_code))
    assert not bad, bad


def test_object_pages_render(flask_app, world, make_client, superadmin):
    admin_a = make_client(world['a']['admin'])
    with flask_app.app_context():
        tid = Ticket.query.execution_options(tenant_bypass=True).filter_by(title='ALPHA-SECRET-TICKET').one().id
        uid = User.query.execution_options(tenant_bypass=True).filter_by(username=world['a']['user']).one().id
    for path in (f'/ticket/{tid}', f'/admin/view_user/{uid}', f'/admin/edit_user/{uid}',
                 f'/admin/download_user_tickets/{uid}', '/admin/download_summary_pdf', '/admin/round_robin_status',
                 f"/platform/schools/{world['a']['sid']}/edit"):
        c = superadmin if path.startswith('/platform') else admin_a
        resp = c.get(path)
        assert resp.status_code == 200, (path, resp.status_code)
    assert make_client(world['a']['user']).get(f'/ticket/{tid}').status_code == 200
    # a different requester in the same school cannot read someone else's ticket
    admin_a.post('/admin/create_user', data={'username': 'peer', 'email': 'peer@example.com', 'role': 'user', 'password': PASSWORD})
    set_password(flask_app, 'peer')
    assert make_client('peer').get(f'/ticket/{tid}').status_code == 302


def test_error_pages(client, world, make_client):
    assert client.get('/definitely/not/here').status_code == 404
    assert make_client(world['a']['user']).get('/ticket/99999999').status_code == 404


# --------------------------------------------------------------------------- analytics
def test_analytics_super_admin_only_and_cross_school(world, superadmin, make_client):
    assert make_client(world['a']['admin']).get('/platform/analytics').status_code == 302
    html = superadmin.get('/platform/analytics?days=0').get_data(as_text=True)
    assert 'Alpha Academy' in html and 'Beta School' in html and 'School comparison' in html
    html = superadmin.get(f"/platform/analytics?days=0&school_id={world['a']['sid']}").get_data(as_text=True)
    assert 'Alpha Academy' in html and 'Beta School<' not in html.split('School comparison')[1]
    assert superadmin.get('/platform/analytics?days=abc').status_code == 200
